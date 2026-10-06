"""
Seyal AI API — Chat Endpoint.

Supports:
- Normal text chat
- Conversation continuity
- File upload
- File content extraction
- AI file analysis
"""

from __future__ import annotations

import asyncio
import json
import uuid
from io import BytesIO
from pathlib import Path
from typing import Any

from fastapi import (
    APIRouter,
    File,
    Form,
    HTTPException,
    Request,
    UploadFile,
)
from fastapi.responses import StreamingResponse

from seyal_ai.api.schemas import (
    ChatResetResponse,
    ChatResponse,
    ToolCallInfo,
)

from seyal_ai.utils.logging import get_logger


log = get_logger("api.chat")

router = APIRouter(
    tags=["chat"]
)


TEXT_EXTENSIONS = {
    ".txt",
    ".md",
    ".py",
    ".js",
    ".ts",
    ".tsx",
    ".jsx",
    ".html",
    ".css",
    ".java",
    ".c",
    ".cpp",
    ".h",
    ".hpp",
    ".json",
    ".xml",
    ".yaml",
    ".yml",
    ".csv",
    ".sql",
}


async def extract_file_content(
    upload: UploadFile,
) -> str:
    """
    Extract readable content from an uploaded file.
    """

    filename = upload.filename or "uploaded_file"

    extension = (
        Path(filename)
        .suffix
        .lower()
    )

    content = await upload.read()

    if not content:
        return (
            f"[File: {filename}]\n"
            "The file is empty."
        )

    # --------------------------------------------------
    # TEXT / CODE FILES
    # --------------------------------------------------

    if extension in TEXT_EXTENSIONS:
        try:
            text = content.decode(
                "utf-8"
            )
        except UnicodeDecodeError:
            text = content.decode(
                "utf-8",
                errors="replace",
            )

        return (
            f"\n--- FILE START: {filename} ---\n"
            f"{text}\n"
            f"--- FILE END: {filename} ---\n"
        )

    # --------------------------------------------------
    # PDF
    # --------------------------------------------------

    if extension == ".pdf":
        try:
            # pyrefly: ignore [missing-import]
            from pypdf import PdfReader

            reader = PdfReader(
                BytesIO(content)
            )

            pages: list[str] = []

            for page in reader.pages:
                page_text = page.extract_text()

                if page_text:
                    pages.append(
                        page_text
                    )

            extracted = "\n\n".join(
                pages
            )

            if not extracted.strip():
                extracted = (
                    "[No readable text could be "
                    "extracted from this PDF.]"
                )

            return (
                f"\n--- PDF START: {filename} ---\n"
                f"{extracted}\n"
                f"--- PDF END: {filename} ---\n"
            )

        except Exception as error:
            raise HTTPException(
                status_code=400,
                detail=(
                    f"Unable to read PDF "
                    f"'{filename}': {error}"
                ),
            )

    # --------------------------------------------------
    # DOCX
    # --------------------------------------------------

    if extension == ".docx":
        try:
            # pyrefly: ignore [missing-import]
            from docx import Document

            document = Document(
                BytesIO(content)
            )

            paragraphs = [
                paragraph.text
                for paragraph in document.paragraphs
                if paragraph.text.strip()
            ]

            extracted = "\n".join(
                paragraphs
            )

            return (
                f"\n--- DOCX START: {filename} ---\n"
                f"{extracted}\n"
                f"--- DOCX END: {filename} ---\n"
            )

        except Exception as error:
            raise HTTPException(
                status_code=400,
                detail=(
                    f"Unable to read DOCX "
                    f"'{filename}': {error}"
                ),
            )

    # --------------------------------------------------
    # EXCEL
    # --------------------------------------------------

    if extension in {
        ".xlsx",
        ".xls",
    }:
        try:
            # pyrefly: ignore [missing-source-for-stubs]
            import openpyxl

            workbook = (
                openpyxl.load_workbook(
                    BytesIO(content),
                    data_only=True,
                )
            )

            output: list[str] = []

            for sheet in workbook.worksheets:

                output.append(
                    f"\nSheet: {sheet.title}"
                )

                for row in sheet.iter_rows(
                    values_only=True
                ):
                    values = [
                        ""
                        if value is None
                        else str(value)
                        for value in row
                    ]

                    output.append(
                        " | ".join(values)
                    )

            extracted = "\n".join(
                output
            )

            return (
                f"\n--- EXCEL START: {filename} ---\n"
                f"{extracted}\n"
                f"--- EXCEL END: {filename} ---\n"
            )

        except Exception as error:
            raise HTTPException(
                status_code=400,
                detail=(
                    f"Unable to read Excel file "
                    f"'{filename}': {error}"
                ),
            )

    # --------------------------------------------------
    # IMAGE
    # --------------------------------------------------

    if upload.content_type and (
        upload.content_type.startswith(
            "image/"
        )
    ):
        try:
            from PIL import Image
            # pyrefly: ignore [missing-import]
            import pytesseract

            image = Image.open(
                BytesIO(content)
            )

            extracted = pytesseract.image_to_string(
                image
            )

            if not extracted.strip():
                extracted = (
                    "[No readable text was detected "
                    "in this image.]"
                )

            return (
                f"\n--- IMAGE OCR START: {filename} ---\n"
                f"{extracted}\n"
                f"--- IMAGE OCR END: {filename} ---\n"
            )

        except Exception as error:
            raise HTTPException(
                status_code=400,
                detail=(
                    f"Unable to analyze image "
                    f"'{filename}': {error}"
                ),
            )

    # --------------------------------------------------
    # UNSUPPORTED
    # --------------------------------------------------

    return (
        f"\n--- FILE: {filename} ---\n"
        f"This file type ({extension or 'unknown'}) "
        f"is not supported for content extraction.\n"
        f"--- END FILE ---\n"
    )


async def _persist_chat_turn(
    target_conv_id: str,
    user_message: str,
    file_names: list[str],
    response_text: str,
    brain: Any | None = None,
) -> None:
    """Non-blocking background helper to persist chat conversation and messages to Supabase database."""
    try:
        from sqlalchemy import select
        from seyal_ai.database.engine import get_session
        from seyal_ai.database.models import Session as DBSession, User
        from seyal_ai.database.repositories.conversation import ConversationRepository

        async with get_session() as session:
            repo = ConversationRepository(session)
            conv = await repo.get_conversation(target_conv_id)

            if conv is None:
                user_res = await session.execute(select(User).limit(1))
                db_user = user_res.scalar_one_or_none()
                if db_user is None:
                    db_user = User(name="User")
                    session.add(db_user)
                    await session.flush()

                sess_res = await session.execute(
                    select(DBSession).where(DBSession.user_id == db_user.id).limit(1)
                )
                db_session = sess_res.scalar_one_or_none()
                if db_session is None:
                    db_session = DBSession(user_id=db_user.id)
                    session.add(db_session)
                    await session.flush()

                summary_source = user_message or (file_names[0] if file_names else "File Analysis")
                # Gather existing titles in simple chat category for strict uniqueness
                all_convs, _ = await repo.list_conversations(offset=0, limit=200)
                existing_titles = {
                    (c.summary or "").strip().lower()
                    for c in all_convs
                    if not (c.summary or "").startswith("[Computer-Use]")
                }
                from seyal_ai.services.title_service import generate_ai_title_for_first_turn
                clean_initial_title = await generate_ai_title_for_first_turn(
                    first_message=summary_source,
                    existing_category_titles=existing_titles,
                    brain=brain,
                    is_computer_use=False,
                )

                conv = await repo.create_conversation(
                    session_id=db_session.id,
                    summary=clean_initial_title,
                    conversation_id=target_conv_id,
                )

            database_user_content = user_message if user_message else "Analyze uploaded file."
            if file_names:
                database_user_content = f"{database_user_content}\n[Attached files: {', '.join(file_names)}]"

            await repo.add_message(
                conversation_id=conv.id,
                role="user",
                content=database_user_content,
            )
            await repo.add_message(
                conversation_id=conv.id,
                role="assistant",
                content=response_text,
            )
            await session.commit()

            # Trigger periodic dynamic title evaluation in background if due
            from seyal_ai.services.title_service import run_periodic_title_update_if_needed
            await run_periodic_title_update_if_needed(
                conversation_id=conv.id,
                session=session,
                repo=repo,
                brain=brain,
                is_computer_use=False,
            )
    except Exception as db_error:
        log.warning("Non-blocking background database persistence notice: %s", db_error)


@router.post(
    "/chat",
    response_model=ChatResponse,
    summary="Send a message to Seyal AI",
    description=(
        "Send a user message with optional "
        "file attachments."
    ),
)
async def chat(
    request: Request,
    message: str = Form(""),
    conversation_id: str | None = Form(None),
    files: list[UploadFile] = File(
        default=[]
    ),
) -> ChatResponse:
    """
    Process a user message and optional files
    through the Seyal AI Brain.
    """

    brain = request.app.state.brain

    # ----------------------------------------------
    # VALIDATION
    # ----------------------------------------------

    if not message.strip() and not files:
        raise HTTPException(
            status_code=400,
            detail=(
                "Please provide a message "
                "or at least one file."
            ),
        )

    # ----------------------------------------------
    # INITIALIZE BRAIN
    # ----------------------------------------------

    if not brain.is_initialized:
        try:
            await brain.initialize()

        except Exception as error:
            log.error(
                "Brain initialization failed: %s",
                error,
            )

            raise HTTPException(
                status_code=503,
                detail=(
                    f"Seyal AI Brain failed "
                    f"to initialize: {error}"
                ),
            ) from error

    # ----------------------------------------------
    # CHECK LLM
    # ----------------------------------------------

    if not brain._router.has_providers:
        raise HTTPException(
            status_code=503,
            detail=(
                "No LLM providers available. "
                "Please configure at least "
                "one API key."
            ),
        )

    # ----------------------------------------------
    # EXTRACT FILE CONTENT
    # ----------------------------------------------

    extracted_files: list[str] = []

    for uploaded_file in files:

        log.info(
            "Analyzing uploaded file: %s",
            uploaded_file.filename,
        )

        extracted = await extract_file_content(
            uploaded_file
        )

        extracted_files.append(
            extracted
        )

    file_context = "\n".join(
        extracted_files
    )

    # ----------------------------------------------
    # BUILD AI INPUT
    # ----------------------------------------------

    user_message = message.strip()

    if file_context:
        if user_message:
            final_input = (
                f"{user_message}\n\n"
                f"Uploaded file content:\n"
                f"{file_context}"
            )
        else:
            final_input = (
                "Analyze the uploaded file "
                "and explain its content.\n\n"
                f"Uploaded file content:\n"
                f"{file_context}"
            )
    else:
        final_input = user_message

    log.info(
        "[BACKEND CHAT REQUEST] "
        "message='%.80s' files=%d",
        user_message,
        len(files),
    )

    # ----------------------------------------------
    # AI PROCESSING
    # ----------------------------------------------

    try:
        response_text = await brain.process(
            final_input,
            allow_tools=True,
        )

    except Exception as error:

        log.error(
            "Brain processing error: %s",
            error,
            exc_info=True,
        )

        raise HTTPException(
            status_code=500,
            detail=(
                f"Error processing message: "
                f"{error}"
            ),
        ) from error

    # ----------------------------------------------
    # NON-BLOCKING BACKGROUND DATABASE PERSISTENCE
    # ----------------------------------------------
    active_conversation_id = conversation_id or str(uuid.uuid4())
    uploaded_filenames = [f.filename or "uploaded_file" for f in files] if files else []

    asyncio.create_task(
        _persist_chat_turn(
            target_conv_id=active_conversation_id,
            user_message=user_message,
            file_names=uploaded_filenames,
            response_text=response_text,
            brain=brain,
        )
    )

    conversation_id = active_conversation_id

    # ----------------------------------------------
    # MODEL INFORMATION
    # ----------------------------------------------
    model_used = None
    try:
        tier = brain._classify_tier(
            final_input
        )

        _, model_name = (
            brain._router
            ._resolve_provider_and_model(
                tier
            )
        )

        model_used = model_name

    except Exception:
        pass

    # ----------------------------------------------
    # TOOL CALL INFORMATION
    # ----------------------------------------------

    tool_calls: list[
        ToolCallInfo
    ] = []

    if (
        hasattr(
            brain,
            "_orchestrator",
        )
        and hasattr(
            brain._orchestrator,
            "last_tool_calls",
        )
    ):

        raw_tools = (
            brain._orchestrator.last_tool_calls
            or []
        )

        for tool in raw_tools:

            if isinstance(
                tool,
                ToolCallInfo,
            ):

                tool_calls.append(
                    tool
                )

            elif isinstance(
                tool,
                dict,
            ):

                tool_calls.append(
                    ToolCallInfo(
                        name=str(
                            tool.get(
                                "name",
                                "unknown",
                            )
                        ),
                        arguments=(
                            tool.get(
                                "arguments",
                                {},
                            )
                            if isinstance(
                                tool.get(
                                    "arguments"
                                ),
                                dict,
                            )
                            else {}
                        ),
                        result=(
                            str(
                                tool.get(
                                    "result"
                                )
                            )
                            if tool.get(
                                "result"
                            )
                            is not None
                            else None
                        ),
                        success=bool(
                            tool.get(
                                "success",
                                True,
                            )
                        ),
                    )
                )

    final_conv_id = (
        conversation_id
        if conversation_id is not None
        else "default"
    )

    return ChatResponse(
        response=response_text,
        conversation_id=final_conv_id,
        model_used=model_used,
        tool_calls=tool_calls,
    )


@router.post(
    "/chat/stream",
    summary="Send a message to Seyal AI with real-time SSE streaming",
    description="Stream response tokens in real-time as they are generated.",
)
async def chat_stream(
    request: Request,
    message: str = Form(""),
    conversation_id: str | None = Form(None),
    files: list[UploadFile] = File(default=[]),
):
    brain = request.app.state.brain

    if not message.strip() and not files:
        raise HTTPException(
            status_code=400,
            detail="Please provide a message or at least one file.",
        )

    if not brain.is_initialized:
        try:
            await brain.initialize()
        except Exception as error:
            log.error("Brain initialization failed: %s", error)
            raise HTTPException(
                status_code=503,
                detail=f"Seyal AI Brain failed to initialize: {error}",
            ) from error

    if not brain._router.has_providers:
        raise HTTPException(
            status_code=503,
            detail="No LLM providers available. Please configure at least one API key.",
        )

    extracted_files: list[str] = []
    for uploaded_file in files:
        extracted = await extract_file_content(uploaded_file)
        extracted_files.append(extracted)

    file_context = "\n".join(extracted_files)
    user_message = message.strip()

    if file_context:
        if user_message:
            final_input = f"{user_message}\n\nUploaded file content:\n{file_context}"
        else:
            final_input = f"Analyze the uploaded file and explain its content.\n\nUploaded file content:\n{file_context}"
    else:
        final_input = user_message

    active_conversation_id = conversation_id or str(uuid.uuid4())
    uploaded_filenames = [f.filename or "uploaded_file" for f in files] if files else []

    async def event_generator():
        yield f"data: {json.dumps({'type': 'start', 'conversation_id': active_conversation_id})}\n\n"
        full_chunks = []
        try:
            from seyal_ai.core.orchestrator import _requires_tools
            if _requires_tools(final_input):
                resp = await brain.process(final_input, allow_tools=True)
                full_chunks.append(resp)
                yield f"data: {json.dumps({'type': 'chunk', 'text': resp})}\n\n"
            else:
                async for chunk in brain.process_stream(final_input, allow_tools=True):
                    full_chunks.append(chunk)
                    yield f"data: {json.dumps({'type': 'chunk', 'text': chunk})}\n\n"
        except Exception as e:
            log.error("Stream generation error: %s", e)
            fallback = "I encountered an error processing your request. Please try again."
            full_chunks.append(fallback)
            yield f"data: {json.dumps({'type': 'chunk', 'text': fallback})}\n\n"

        response_text = "".join(full_chunks)
        yield f"data: {json.dumps({'type': 'done', 'conversation_id': active_conversation_id, 'full_response': response_text})}\n\n"

        # Background persistence & title service
        asyncio.create_task(
            _persist_chat_turn(
                target_conv_id=active_conversation_id,
                user_message=user_message,
                file_names=uploaded_filenames,
                response_text=response_text,
                brain=brain,
            )
        )

    return StreamingResponse(
        event_generator(),
        media_type="text/event-stream",
        headers={
            "Cache-Control": "no-cache",
            "Connection": "keep-alive",
            "X-Accel-Buffering": "no",
        },
    )


@router.post(
    "/chat/reset",
    response_model=ChatResetResponse,
    summary="Reset conversation",
    description=(
        "Clear the current conversation "
        "context and start fresh."
    ),
)
async def reset_chat(
    request: Request,
) -> ChatResetResponse:
    """Reset the conversation context."""

    brain = request.app.state.brain

    if brain.is_initialized:
        brain.reset_conversation()

    log.info(
        "Conversation reset via API"
    )

    return ChatResetResponse(
        message=(
            "Conversation reset successfully."
        )
    )